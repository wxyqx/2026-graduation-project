"""
【这个文件是干什么的？】
把投递记录变成一张 Excel 表（或 csv 文件）让用户下载。

三步：查出符合条件的投递 → 按用户勾选的列整理成表格（表头中文、状态翻译成中文）→ 生成文件字节。
生成 Excel 用的工具叫 openpyxl。
"""
import csv
import io
from datetime import datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Application
from app.schemas.export import ExportFilters
from app.services.state_machine import STAGE_LABELS

# 所有可导出的列：程序里的名字 → Excel 表头的中文。字典的顺序就是导出时列的顺序
EXPORT_FIELDS: dict[str, str] = {
    "id": "投递编号",
    "candidate_name": "候选人姓名",
    "position_name": "岗位名称",
    "current_stage": "当前阶段",
    "overall_status": "全局状态",
    "ai_result": "AI筛选结果",
    "ai_comment": "AI筛选理由",
    "resume_submit_time": "简历提交时间",
    "resume_result": "简历筛选结果",
    "phone_time": "电话沟通时间",
    "phone_result": "电话沟通结果",
    "test_time": "笔试时间",
    "test_result": "笔试结果",
    "pro_time": "专业面时间",
    "pro_result": "专业面结果",
    "hr_time": "HR面时间",
    "hr_result": "HR面结果",
    "final_time": "终面时间",
    "final_result": "终面结果",
    "create_time": "投递创建时间",
    "update_time": "更新时间",
}

# 数据库里存的是英文单词，导出时翻译成中文，用户看得懂
STATUS_LABELS = {"pending": "进行中", "pass": "已录用", "fail": "已淘汰"}
RESULT_LABELS = {"pass": "通过", "fail": "淘汰"}


def query_applications(db: Session, f: ExportFilters) -> list[Application]:
    """按筛选条件查投递。没填的条件就不加限制。"""
    stmt = select(Application).order_by(Application.id)
    if f.start_date:
        # 开始日期那天的 0 点起
        stmt = stmt.where(Application.create_time >= datetime.combine(f.start_date, time.min))
    if f.end_date:
        # 结束日期的「第二天 0 点」之前 = 包含结束日期整天
        stmt = stmt.where(Application.create_time < datetime.combine(f.end_date + timedelta(days=1), time.min))
    if f.pos_id:
        stmt = stmt.where(Application.pos_id == f.pos_id)
    if f.stage:
        stmt = stmt.where(Application.current_stage == f.stage)
    if f.status:
        stmt = stmt.where(Application.overall_status == f.status)
    return db.scalars(stmt).all()


def _cell(app: Application, key: str):
    """算出某一条投递、某一列该填什么值（一个单元格）。顺便把英文翻成中文、把时间格式化。"""
    if key == "candidate_name":
        return app.candidate.name if app.candidate else None
    if key == "position_name":
        return app.position.position_name if app.position else None
    val = getattr(app, key)  # 其余列直接按名字从投递对象上取
    if key == "current_stage":
        return STAGE_LABELS.get(val, val)  # .get(val, val)：查得到就翻译，查不到就原样返回
    if key == "overall_status":
        return STATUS_LABELS.get(val, val)
    if key.endswith("_result"):
        return RESULT_LABELS.get(val, val)
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")  # 时间统一成「2026-09-08 14:30:00」这种格式
    return val


def build_rows(apps: list[Application], fields: list[str]) -> tuple[list[str], list[list]]:
    """整理成「表头 + 很多行」。表头是中文列名，每一行是一条投递按列顺序排好的值。"""
    headers = [EXPORT_FIELDS[k] for k in fields]
    rows = [[_cell(a, k) for k in fields] for a in apps]
    return headers, rows


def to_xlsx(headers: list[str], rows: list[list], title: str = "投递记录") -> bytes:
    """生成 Excel 文件，返回文件的字节（不落盘，直接发给浏览器下载）。

    title：工作表标签名。导出交叉表时会传「阶段岗位汇总」。
    """
    wb = Workbook()  # 新建一个工作簿
    ws = wb.active  # 拿到默认的第一张工作表
    ws.title = title[:31]  # 工作表名最多 31 个字符
    ws.append(headers)  # 第一行写表头
    for r in rows:
        ws.append(r)  # 之后每行写一条
    # 根据每列内容的长度自动调列宽，免得打开时全是「####」或挤在一起
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [8])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width * 1.6, 10), 60)  # 中文占两格，乘 1.6；限制在 10～60
    ws.freeze_panes = "A2"  # 冻结第一行，往下滚时表头一直可见
    buf = io.BytesIO()  # 内存里的「假文件」
    wb.save(buf)
    return buf.getvalue()


def to_csv(headers: list[str], rows: list[list]) -> bytes:
    """生成 csv（纯文本表格）。用 utf-8-sig 编码——开头带一个特殊标记，Excel 打开才不会把中文显示成乱码。"""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])  # 空值写成空字符串，而不是 "None"
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx_sheets(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    """生成含多个工作表的 Excel。

    sheets：[(工作表名, 表头, 数据行), ...]
    用于「招聘周报」：第一张放阶段岗位汇总，第二张放进行中候选人清单。
    """
    wb = Workbook()
    # 用第一张替换掉默认的空表
    for idx, (name, headers, rows) in enumerate(sheets):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = (name or f"表{idx + 1}")[:31]  # 工作表名上限 31 字符
        ws.append(headers)
        for r in rows:
            ws.append(r)
        # 自动列宽（中文按 1.6 倍估算）
        for i, h in enumerate(headers, start=1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [8])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width * 1.6, 10), 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
